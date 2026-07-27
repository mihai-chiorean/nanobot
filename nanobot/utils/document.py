"""Document text extraction utilities for nanobot."""

import mimetypes
from pathlib import Path

from loguru import logger

from nanobot.utils.helpers import detect_image_mime

# Supported file extensions for text extraction
SUPPORTED_EXTENSIONS: set[str] = {
    # Document formats
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    # Text formats
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    # Image formats (for future OCR support)
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
}

_MAX_TEXT_LENGTH = 200_000
_MAX_PDF_PAGES = 100
_MAX_DOCX_PARAGRAPHS = 20_000
_MAX_XLSX_SHEETS = 50
_MAX_XLSX_ROWS = 50_000
_MAX_XLSX_CELLS = 500_000
_MAX_PPTX_SLIDES = 500
_MAX_PPTX_SHAPES = 20_000


class _TextBudget:
    """Collect bounded extracted text without retaining discarded content."""

    def __init__(self, max_length: int = _MAX_TEXT_LENGTH):
        self._max_length = max_length
        self._parts: list[str] = []
        self._length = 0
        self.truncated = False

    @property
    def exhausted(self) -> bool:
        return self._length >= self._max_length

    def append(self, value: str, *, separator: str = "") -> None:
        if self.exhausted:
            self.truncated = True
            return
        combined = f"{separator}{value}" if self._parts else value
        remaining = self._max_length - self._length
        if len(combined) > remaining:
            self._parts.append(combined[:remaining])
            self._length = self._max_length
            self.truncated = True
            return
        self._parts.append(combined)
        self._length += len(combined)

    def render(self) -> str:
        result = "".join(self._parts)
        if self.truncated:
            return result + "... (truncated at extraction limit)"
        return result


def extract_text(path: Path) -> str | None:
    """Extract text from a file.

    Args:
        path: Path to the file.

    Returns:
        Extracted text as string, None for unsupported types,
        or error string for failures.
    """
    if not isinstance(path, Path):
        path = Path(path)

    if not path.exists():
        return f"[error: file not found: {path}]"

    ext = path.suffix.lower()

    # Document formats -- each branch lazily imports its parser so that
    # startup does not pay the ~25 MB cost of loading openpyxl /
    # python-docx / python-pptx / pypdf up front (see issue #3422).
    if ext == ".pdf":
        return _extract_pdf(path)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext == ".xlsx":
        return _extract_xlsx(path)
    elif ext == ".pptx":
        return _extract_pptx(path)
    elif _is_text_extension(ext):
        return _extract_text_file(path)
    elif ext in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        # Image files - for future OCR support
        return f"[image: {path.name}]"
    else:
        # Unsupported extension
        return None


def _extract_pdf(path: Path) -> str:
    """Extract text from PDF using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[error: pypdf not installed]"
    try:
        reader = PdfReader(path)
        budget = _TextBudget()
        page_count = len(reader.pages)
        for i, page in enumerate(reader.pages[:_MAX_PDF_PAGES], 1):
            text = page.extract_text() or ""
            budget.append(f"--- Page {i} ---\n{text}", separator="\n\n")
            if budget.exhausted:
                break
        if page_count > _MAX_PDF_PAGES:
            budget.truncated = True
        return budget.render()
    except Exception as e:
        logger.error("Failed to extract PDF {}: {}", path, e)
        return f"[error: failed to extract PDF: {e!s}]"


def _extract_docx(path: Path) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return "[error: python-docx not installed]"
    try:
        doc = DocxDocument(path)
        budget = _TextBudget()
        for index, paragraph in enumerate(doc.paragraphs):
            if index >= _MAX_DOCX_PARAGRAPHS:
                budget.truncated = True
                break
            if paragraph.text.strip():
                budget.append(paragraph.text, separator="\n\n")
                if budget.exhausted:
                    break
        return budget.render()
    except Exception as e:
        logger.error("Failed to extract DOCX {}: {}", path, e)
        return f"[error: failed to extract DOCX: {e!s}]"


def _extract_xlsx(path: Path) -> str:
    """Extract text from XLSX using openpyxl."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[error: openpyxl not installed]"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            budget = _TextBudget()
            row_count = 0
            cell_count = 0
            for sheet_index, sheet_name in enumerate(wb.sheetnames):
                if sheet_index >= _MAX_XLSX_SHEETS:
                    budget.truncated = True
                    break
                ws = wb[sheet_name]
                wrote_sheet_header = False
                for row in ws.iter_rows(values_only=True):
                    row_count += 1
                    cell_count += len(row)
                    if row_count > _MAX_XLSX_ROWS or cell_count > _MAX_XLSX_CELLS:
                        budget.truncated = True
                        break
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        if not wrote_sheet_header:
                            budget.append(f"--- Sheet: {sheet_name} ---", separator="\n\n")
                            wrote_sheet_header = True
                        budget.append(row_text, separator="\n")
                        if budget.exhausted:
                            break
                if budget.truncated or budget.exhausted:
                    break
            return budget.render()
        finally:
            wb.close()
    except Exception as e:
        logger.error("Failed to extract XLSX {}: {}", path, e)
        return f"[error: failed to extract XLSX: {e!s}]"


def _extract_pptx(path: Path) -> str:
    """Extract text from PPTX using python-pptx."""
    try:
        from pptx import Presentation as PptxPresentation
    except ImportError:
        return "[error: python-pptx not installed]"
    try:
        prs = PptxPresentation(path)
        budget = _TextBudget()
        shape_count = 0
        slide_count = len(prs.slides)
        for i, slide in enumerate(prs.slides[:_MAX_PPTX_SLIDES], 1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                shape_count = _collect_pptx_shape_text(
                    shape,
                    slide_text,
                    current_count=shape_count,
                    max_shapes=_MAX_PPTX_SHAPES,
                )
                if shape_count >= _MAX_PPTX_SHAPES:
                    budget.truncated = True
                    break
            if slide_text:
                budget.append(f"--- Slide {i} ---\n" + "\n".join(slide_text), separator="\n\n")
            if budget.truncated or budget.exhausted:
                break
        if slide_count > _MAX_PPTX_SLIDES:
            budget.truncated = True
        return budget.render()
    except Exception as e:
        logger.error("Failed to extract PPTX {}: {}", path, e)
        return f"[error: failed to extract PPTX: {e!s}]"


def _collect_pptx_shape_text(
    shape,
    out: list[str],
    *,
    current_count: int = 0,
    max_shapes: int = _MAX_PPTX_SHAPES,
) -> int:
    """Collect text from a PPTX shape, recursing into groups and tables.

    Groups have ``has_text_frame=False`` and must be walked via ``.shapes``;
    tables are GraphicFrame objects whose cell text lives under ``.table``.
    """
    if current_count >= max_shapes:
        return current_count
    current_count += 1

    sub_shapes = getattr(shape, "shapes", None)
    if sub_shapes is not None:
        for sub in sub_shapes:
            current_count = _collect_pptx_shape_text(
                sub,
                out,
                current_count=current_count,
                max_shapes=max_shapes,
            )
            if current_count >= max_shapes:
                break
        return current_count

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(cell for cell in cells if cell)
            if line:
                out.append(line)
        return current_count

    text = getattr(shape, "text", "")
    if text:
        out.append(text)
    return current_count


def _extract_text_file(path: Path) -> str:
    """Extract text from a plain text file."""
    try:
        # Try UTF-8 first, then latin-1 fallback
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="latin-1")
        return _truncate(content, _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.error("Failed to read text file {}: {}", path, e)
        return f"[error: failed to read file: {e!s}]"


def _truncate(text: str, max_length: int) -> str:
    """Truncate text with a suffix indicating truncation."""
    if len(text) <= max_length:
        return text
    return text[:max_length] + f"... (truncated, {len(text)} chars total)"


def _is_text_extension(ext: str) -> bool:
    """Check if extension is a text format."""
    return ext in {
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".xml",
        ".html",
        ".htm",
        ".log",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
    }


# ---------------------------------------------------------------------------
# High-level helper: split media into images + extracted document text
# ---------------------------------------------------------------------------

_MAX_EXTRACT_FILE_SIZE = 10 * 1024 * 1024


def extract_documents(
    text: str,
    media_paths: list[str],
    *,
    max_file_size: int = _MAX_EXTRACT_FILE_SIZE,
) -> tuple[str, list[str]]:
    """Separate images from documents in *media_paths*.

    Documents (PDF, DOCX, XLSX, PPTX, plain-text, …) have their text
    extracted and appended to *text*.  Only image paths are kept in the
    returned list so that downstream layers only need to handle vision
    blocks.

    Files larger than *max_file_size* bytes are skipped with a warning
    to avoid unbounded memory / CPU usage.
    """
    image_paths: list[str] = []
    doc_texts: list[str] = []

    for path_str in media_paths:
        p = Path(path_str)
        if not p.is_file():
            continue

        try:
            size = p.stat().st_size
        except OSError:
            continue
        if size > max_file_size:
            logger.warning(
                "Skipping oversized file for extraction: {} ({:.1f} MB > {} MB limit)",
                p.name, size / (1024 * 1024), max_file_size // (1024 * 1024),
            )
            continue

        with open(p, "rb") as f:
            header = f.read(16)
        mime = detect_image_mime(header) or mimetypes.guess_type(path_str)[0]
        if mime and mime.startswith("image/"):
            image_paths.append(path_str)
        else:
            extracted = extract_text(p)
            if extracted and not extracted.startswith("[error:"):
                doc_texts.append(f"[File: {p.name}]\n{extracted}")

    if doc_texts:
        text = text + "\n\n" + "\n\n".join(doc_texts)

    return text, image_paths
